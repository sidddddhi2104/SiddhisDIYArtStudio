from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# Create workbook
wb = Workbook()
ws = wb.active
ws.title = "Test Cases"

# Headers
headers = [
    "Test Case ID",
    "Module",
    "Feature",
    "Test Scenario",
    "Preconditions",
    "Test Steps",
    "Test Data",
    "Expected Result",
    "Actual Result",
    "Status",
    "Priority",
    "Tester",
    "Date",
    "Remarks"
]

# Header Style
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)

for col, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col)
    cell.value = header
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Sample Test Cases
sample_data = [
    [
        "TC001",
        "Login",
        "User Login",
        "Verify login with valid credentials",
        "User account exists",
        "1. Open app\n2. Enter email\n3. Enter password\n4. Click Login",
        "Email: user@test.com\nPassword: Password123",
        "User should be redirected to Dashboard",
        "",
        "",
        "High",
        "",
        "",
        ""
    ],
    [
        "TC002",
        "Login",
        "User Login",
        "Verify login with invalid password",
        "User account exists",
        "1. Open app\n2. Enter email\n3. Enter wrong password\n4. Click Login",
        "Email: user@test.com\nPassword: Wrong123",
        "Error message should be displayed",
        "",
        "",
        "High",
        "",
        "",
        ""
    ]
]

# Insert Sample Data
for row in sample_data:
    ws.append(row)

# Set Column Widths
column_widths = {
    "A": 15,
    "B": 20,
    "C": 25,
    "D": 35,
    "E": 25,
    "F": 45,
    "G": 30,
    "H": 35,
    "I": 35,
    "J": 15,
    "K": 12,
    "L": 20,
    "M": 15,
    "N": 25,
}

for col, width in column_widths.items():
    ws.column_dimensions[col].width = width

# Freeze Header
ws.freeze_panes = "A2"

# Save Excel File
file_name = "Sivika_Test_Cases.xlsx"
wb.save(file_name)

print(f"Excel file '{file_name}' created successfully.")